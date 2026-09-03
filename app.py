import streamlit as st
import pandas as pd
import io
import os
import sqlite3
import re
from dotenv import load_dotenv

from src.modules.m1_teacher.parser import extract_raw_from_excel, process_hybrid_ai
from src.modules.m1_teacher.database import save_teachers_to_db, load_teachers_from_db, init_db as init_db_m1
from src.modules.m2_facility.parser import extract_lc_cd_matrix, extract_all_pinned_from_excel, extract_assignments_matrix
from src.modules.m2_facility.database import init_facility_db, seed_default_pinned_slots, load_table
from src.modules.m3_validator.checker import get_fixed_schedule
from src.modules.m4_engine.solver import run_round_1, run_round_2, run_round_3

load_dotenv(override=True)
st.set_page_config(page_title="TKB-AI 1.0", page_icon="🏫", layout="wide")

init_db_m1()
init_facility_db()
seed_default_pinned_slots()

DB_PATH = os.path.join("data", "database", "tkb_data.db")

def save_table_to_db(df: pd.DataFrame, table_name: str):
    conn = sqlite3.connect(DB_PATH)
    df.to_sql(table_name, conn, if_exists="replace", index=False)
    conn.close()

# --- BẢNG NGUYÊN TẮC THỰC TẾ ---
DEFAULT_RULES = [
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "MAX_2_TIET", "Mô tả": "Không xếp 1 môn quá 2 tiết/buổi/lớp (trừ ghép).", "Loại": "Cơ bản"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "NO_CACH_TIET", "Mô tả": "Tránh xếp cách tiết (Tránh 1-3, 2-4, 1-4, 5-7).", "Loại": "Cơ bản"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "TEACHER_PREFS", "Mô tả": "Áp dụng ngày nghỉ cá nhân và nghỉ theo Tổ bộ môn.", "Loại": "Cơ bản"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "BLOCK_MON_CHINH", "Mô tả": "Các môn cốt lõi (Toán, Văn, Anh...) ép xếp liền 2 tiết.", "Loại": "Nâng cao"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "MIN_2_TIET_GV", "Mô tả": "Tối ưu lịch GV: Tránh 1 buổi chỉ dạy đúng 1 tiết.", "Loại": "Nâng cao"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "NO_AFTERNOON_GAP", "Mô tả": "Chống thủng lỗ tiết 6 buổi chiều (mô hình 1-0-1).", "Loại": "Nâng cao"},
    {"Bật/Tắt": True, "Mã Nguyên Tắc": "KIN_B_SANG", "Mô tả": "Ép thuật toán ưu tiên lấp kín toàn bộ buổi sáng.", "Loại": "Nâng cao"}
]

def load_rules_from_db():
    try:
        conn = sqlite3.connect(DB_PATH)
        df = pd.read_sql("SELECT * FROM rules_config", conn)
        df['Bật/Tắt'] = df['Bật/Tắt'].astype(bool)
        conn.close()
        return df
    except:
        return pd.DataFrame(DEFAULT_RULES)

def save_rules_to_db_persist(df):
    conn = sqlite3.connect(DB_PATH)
    df_save = df.copy()
    df_save['Bật/Tắt'] = df_save['Bật/Tắt'].astype(int)
    df_save.to_sql("rules_config", conn, if_exists="replace", index=False)
    conn.close()

# ==============================================================
# HÀM XUẤT EXCEL TKB
# ==============================================================
def generate_styled_excel_both_shifts(df_sang, df_chieu, df_raw):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        out_buffer = io.BytesIO()
        with pd.ExcelWriter(out_buffer, engine="xlsxwriter") as writer:
            df_sang.to_excel(writer, sheet_name="TKB Sáng", index=False)
            df_chieu.to_excel(writer, sheet_name="TKB Chiều", index=False)
        return out_buffer.getvalue()

    wb = Workbook()
    
    fill_tnhn = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_tnhn = Font(color="1F4E78", bold=True)
    fill_nhom = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    font_nhom = Font(color="B08000", bold=True)
    fill_do = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    font_do = Font(color="C00000", bold=True)
    fill_xanh = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_xanh = Font(color="375623", bold=True)
    
    font_bold = Font(bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    def write_styled_sheet(ws, df, sheet_title):
        ws.title = sheet_title
        if df.empty: return
        
        pt = df.pivot_table(index=['Ngày', 'Tiết'], columns='Lớp', values='Giáo viên', aggfunc=lambda x: ' / '.join(x)).fillna('')
        def sort_class(c):
            m = re.match(r'^(\d+)A(\d+)$', c)
            if m: return (int(m.group(1)), int(m.group(2)))
            return (99, 99)
        pt = pt.reindex(sorted(pt.columns, key=sort_class), axis=1)
        
        day_order = {"Thứ Hai": 1, "Thứ Ba": 2, "Thứ Tư": 3, "Thứ Năm": 4, "Thứ Sáu": 5}
        pt = pt.reindex(sorted(pt.index, key=lambda x: (day_order.get(x[0], 99), x[1])))

        ws.cell(row=1, column=1, value="Ngày").font = font_bold
        ws.cell(row=1, column=1).alignment = align_center
        ws.cell(row=1, column=1).border = border_thin
        ws.column_dimensions['A'].width = 12

        ws.cell(row=1, column=2, value="Tiết").font = font_bold
        ws.cell(row=1, column=2).alignment = align_center
        ws.cell(row=1, column=2).border = border_thin
        ws.column_dimensions['B'].width = 8

        for col_idx, c in enumerate(pt.columns):
            cell = ws.cell(row=1, column=col_idx+3, value=c)
            cell.font = font_bold
            cell.alignment = align_center
            cell.border = border_thin
            ws.column_dimensions[get_column_letter(col_idx+3)].width = 18

        row_idx = 2
        current_day = None
        start_row_merge = 2

        for (d, p), row_data in pt.iterrows():
            if d != current_day:
                if current_day is not None and (row_idx - 1) >= start_row_merge:
                    try: ws.merge_cells(start_row=start_row_merge, start_column=1, end_row=row_idx-1, end_column=1)
                    except: pass
                current_day = d
                start_row_merge = row_idx
                c_day = ws.cell(row=row_idx, column=1, value=d)
                c_day.font = font_bold
                c_day.alignment = align_center
            else:
                ws.cell(row=row_idx, column=1).border = border_thin
            
            ws.cell(row=row_idx, column=1).border = border_thin
            c_period = ws.cell(row=row_idx, column=2, value=p)
            c_period.font = font_bold
            c_period.alignment = align_center
            c_period.border = border_thin
            
            for col_idx, c in enumerate(pt.columns):
                val = row_data[c]
                cell = ws.cell(row=row_idx, column=col_idx+3, value=val)
                cell.alignment = align_center
                cell.border = border_thin
                
                val_str = str(val)
                if "TNHN" in val_str:
                    cell.fill = fill_tnhn
                    cell.font = font_tnhn
                elif "⭐" in val_str or "Nhóm" in val_str or "Lớp ghép" in val_str:
                    cell.fill = fill_nhom
                    cell.font = font_nhom
                elif "🔴" in val_str or "⛔" in val_str:
                    cell.fill = fill_do
                    cell.font = font_do
                elif "🟢" in val_str:
                    cell.fill = fill_xanh
                    cell.font = font_xanh
                elif "📌" in val_str:
                    cell.font = font_bold
            row_idx += 1
            
        if current_day is not None and (row_idx - 1) >= start_row_merge:
            try: ws.merge_cells(start_row=start_row_merge, start_column=1, end_row=row_idx-1, end_column=1)
            except: pass

    ws_sang = wb.active
    write_styled_sheet(ws_sang, df_sang, "TKB Sáng")
    ws_chieu = wb.create_sheet(title="TKB Chiều")
    write_styled_sheet(ws_chieu, df_chieu, "TKB Chiều")
    
    ws_raw = wb.create_sheet(title="Dữ Liệu Thô")
    from openpyxl.utils.dataframe import dataframe_to_rows
    for r in dataframe_to_rows(df_raw, index=False, header=True):
        ws_raw.append(r)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_single_styled_excel(df_pivot, entity_name, entity_type="Lớp"):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        out_buffer = io.BytesIO()
        with pd.ExcelWriter(out_buffer, engine="xlsxwriter") as writer:
            df_pivot.to_excel(writer, sheet_name=f"TKB {entity_name}")
        return out_buffer.getvalue()

    wb = Workbook()
    ws = wb.active
    ws.title = f"TKB {entity_name}"

    font_bold = Font(bold=True)
    font_title = Font(bold=True, size=14, color="1E3A8A")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_tnhn = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_nhom = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    fill_do = PatternFill(start_color="FCA5A5", end_color="FCA5A5", fill_type="solid")
    fill_xanh = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    ws.merge_cells('A1:F1')
    ws['A1'] = f"THỜI KHÓA BIỂU {entity_type.upper()}: {entity_name}"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center

    headers = ["Tiết"] + list(df_pivot.columns)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_bold
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin
    
    ws.column_dimensions['A'].width = 8
    for col_idx in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    row_idx = 4
    for tiet, row in df_pivot.iterrows():
        c = ws.cell(row=row_idx, column=1, value=tiet)
        c.font = font_bold
        c.alignment = align_center
        c.border = border_thin
        
        for col_idx, day in enumerate(df_pivot.columns, 2):
            val = str(row[day])
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.alignment = align_center
            c.border = border_thin
            
            if "TNHN" in val:
                c.fill = fill_tnhn
                c.font = Font(color="1E40AF", bold=True)
            elif "⭐" in val or "Nhóm" in val or "Lớp ghép" in val:
                c.fill = fill_nhom
                c.font = Font(color="854D0E", bold=True)
            elif "🔴" in val or "⛔" in val or "❌" in val:
                c.fill = fill_do
                c.font = Font(color="7F1D1D", bold=True)
            elif "🟢" in val:
                c.fill = fill_xanh
                c.font = Font(color="166534", bold=True)
            elif "📌" in val:
                c.font = font_bold
        row_idx += 1

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ==============================================================
# GIAO DIỆN & CÁC HÀM TIỆN ÍCH STREAMLIT
# ==============================================================
st.markdown("""
<style>
    .main-title { font-size: 26px; font-weight: bold; color: #1E3A8A; margin-bottom: 5px; }
    .stButton>button { border-radius: 6px; font-weight: 600; }
</style>
""", unsafe_allow_html=True)

st.sidebar.title("🏫 TKB-AI-1.0")
menu = st.sidebar.radio("Điều hướng Module", ["Module 1: Giáo viên", "Module 2: Lớp & CSVC", "Module 3: Kiểm tra Lịch", "Module 4: AI Lấp Đầy"])

st.sidebar.markdown("---")
if st.sidebar.button("🗑️ Xóa sạch Dữ liệu (Reset)"):
    if os.path.exists(DB_PATH):
        try: os.remove(DB_PATH)
        except Exception as e: st.sidebar.error(f"Lỗi: {e}")
    for key in list(st.session_state.keys()): del st.session_state[key]
    st.rerun()

st.sidebar.markdown("### 💾 QUẢN LÝ DỮ LIỆU (CLOUD)")
st.sidebar.info("Sử dụng tính năng này để tránh mất dữ liệu khi máy chủ Cloud khởi động lại.")

if os.path.exists(DB_PATH):
    with open(DB_PATH, "rb") as f:
        st.sidebar.download_button(
            label="📥 Tải Backup DB về máy",
            data=f,
            file_name="tkb_data_backup.db",
            mime="application/octet-stream",
            use_container_width=True
        )

uploaded_db = st.sidebar.file_uploader("📤 Phục hồi DB từ máy", type=["db"])
if uploaded_db is not None:
    if st.sidebar.button("Tiến hành Phục hồi", type="primary", use_container_width=True):
        os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
        with open(DB_PATH, "wb") as f:
            f.write(uploaded_db.getbuffer())
        st.sidebar.success("✅ Phục hồi thành công!")
        st.rerun()
st.sidebar.markdown("---")

def pivot_main(df):
    try: 
        pt = df.pivot_table(index=['Ngày', 'Tiết'], columns='Lớp', values='Giáo viên', aggfunc=lambda x: ' / '.join(x)).fillna('')
        def sort_class(c):
            m = re.match(r'^(\d+)A(\d+)$', c)
            if m: return (int(m.group(1)), int(m.group(2)))
            return (99, 99)
        pt = pt.reindex(sorted(pt.columns, key=sort_class), axis=1)
        day_order = {"Thứ Hai": 1, "Thứ Ba": 2, "Thứ Tư": 3, "Thứ Năm": 4, "Thứ Sáu": 5}
        pt = pt.reindex(sorted(pt.index, key=lambda x: (day_order.get(x[0], 99), x[1])))
        return pt
    except: 
        return pd.DataFrame()

def pivot_single(df, val_col):
    try: 
        pt = df.pivot_table(index='Tiết', columns='Ngày', values=val_col, aggfunc=lambda x: ' / '.join(x)).fillna('')
        for d in ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"]:
            if d not in pt.columns: pt[d] = ""
        return pt[["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"]]
    except: return pd.DataFrame()

def styler_func(df):
    def color_cells(val):
        if not isinstance(val, str): return ''
        if '⛔' in val: return 'background-color: #fee2e2; color: #991b1b; font-weight: bold; text-decoration: line-through;'
        if 'TNHN' in val: return 'background-color: #dbeafe; color: #1e40af; font-weight: bold;'
        if '⭐' in val: return 'background-color: #fef08a; color: #854d0e; font-weight: bold;'
        if '🟢' in val: return 'background-color: #dcfce7; color: #166534; font-weight: bold;'
        if '📌' in val: return 'background-color: #f3e8ff; color: #7e22ce; font-weight: bold;'
        if '🔴' in val: return 'background-color: #fca5a5; color: #7f1d1d; font-weight: bold;' 
        if val == '': return 'background-color: #f3f4f6;'
        return 'background-color: white; color: black;'
    return df.style.map(color_cells) if hasattr(df.style, 'map') else df.style.applymap(color_cells)

# --- HÀM HỖ TRỢ XỬ LÝ DỮ LIỆU TỪ FILE CHỈNH SỬA ---
def parse_r2_vars_from_df(new_df_r2):
    new_r2_vars = []
    try: teachers_df = load_teachers_from_db()
    except: teachers_df = pd.DataFrame()
    
    for _, r in new_df_r2.iterrows():
        c = str(r.get('Lớp', '')).strip()
        gv_str = str(r.get('Giáo viên', '')).strip()
        d = str(r.get('Ngày', '')).strip()
        try: p = int(r.get('Tiết', 0))
        except: continue
        
        if not c or not d or p == 0: continue
        if any(x in gv_str for x in ['⭐', '🟢']): continue
        
        clean_gv_str = gv_str.replace('🔴', '').replace('📌', '').strip()
        
        if '-' in clean_gv_str:
            parts = clean_gv_str.split('-', 1)
            mon = parts[0].strip()
            gv_ho_ten = parts[1].strip()
            new_r2_vars.append((c, mon, gv_ho_ten, d, p))
        elif clean_gv_str == "TNHN":
            for _, t_row in teachers_df.iterrows():
                if str(t_row.get('lop_chu_nhiem', '')).strip() == c:
                    new_r2_vars.append((c, "TNHN", str(t_row.get('ho_ten', '')).strip(), d, p))
                    break
    return new_r2_vars


if menu == "Module 1: Giáo viên":
    st.markdown('<div class="main-title">🎯 MODULE 1: QUẢN TRỊ DỮ LIỆU GIÁO VIÊN</div>', unsafe_allow_html=True)
    if "teachers_df" not in st.session_state: st.session_state.teachers_df = load_teachers_from_db()
    if "classes_totals_df" not in st.session_state: st.session_state.classes_totals_df = pd.DataFrame()

    with st.sidebar:
        uploaded_file_m1 = st.file_uploader("Nạp Excel phân công", type=["xls", "xlsx"], key="file_m1")

    if uploaded_file_m1 and st.button("⚡ Bóc tách & Chuẩn hóa", use_container_width=True, type="primary"):
        with st.spinner("Đang bóc tách môn học và phân tích tổng tiết..."):
            df_gv_raw, df_classes = extract_raw_from_excel(uploaded_file_m1)
            st.session_state.teachers_df = pd.DataFrame(process_hybrid_ai(df_gv_raw))
            st.session_state.classes_totals_df = df_classes
            st.success("✅ Đã chuẩn hóa nhanh thành công!")

    if not st.session_state.teachers_df.empty:
        tab_gv, tab_lop = st.tabs(["👩‍🏫 Danh sách Giáo Viên", "📋 Tổng Tiết Các Lớp"])
        with tab_gv:
            t_col1, _, t_col3 = st.columns([3, 4, 3])
            with t_col1: selected_to = st.selectbox("Lọc Tổ:", ["Tất cả"] + sorted(list(st.session_state.teachers_df["to_bo_mon"].dropna().unique())))
            with t_col3: save_btn_m1 = st.button("💾 LƯU DATABASE", type="primary", use_container_width=True)

            filtered_df = st.session_state.teachers_df if selected_to == "Tất cả" else st.session_state.teachers_df[st.session_state.teachers_df["to_bo_mon"] == selected_to]
            edited_df = st.data_editor(filtered_df, use_container_width=True, num_rows="dynamic", height=550)
            
            if selected_to == "Tất cả": st.session_state.teachers_df = edited_df
            else: st.session_state.teachers_df.update(edited_df)
            if save_btn_m1: save_teachers_to_db(st.session_state.teachers_df)
            
        with tab_lop:
            st.info("Bảng đối chiếu tổng số tiết thực tế của từng lớp so với định mức khối.")
            if not st.session_state.classes_totals_df.empty:
                st.dataframe(st.session_state.classes_totals_df, use_container_width=True, height=550)

elif menu == "Module 2: Lớp & CSVC":
    st.markdown('<div class="main-title">🏢 MODULE 2: QUẢN LÝ LỚP & CƠ SỞ VẬT CHẤT</div>', unsafe_allow_html=True)
    tab1, tab2 = st.tabs(["🧩 Ma Trận Lớp", "🔒 Tiết Cố định"])
    
    with st.sidebar: uploaded_file_m2 = st.file_uploader("Nạp Excel tổng hợp", type=["xls", "xlsx"], key="file_m2")

    with tab1:
        if uploaded_file_m2 and st.button("Trích xuất Cấu trúc Lớp", type="primary"):
            df_cd, df_lc = extract_lc_cd_matrix(uploaded_file_m2)
            save_table_to_db(df_cd, "lop_chuyen_de")
            save_table_to_db(df_lc, "nhom_lua_chon")
            st.success("✅ Thành công!")
                    
    with tab2:
        if uploaded_file_m2 and st.button("🔍 Quét TOÀN BỘ Dữ Liệu TKB Cố Định", type="secondary"):
            with st.spinner("Đang quét động mọi ô trống..."):
                df_assign = extract_assignments_matrix(uploaded_file_m2) 
                if not df_assign.empty: save_table_to_db(df_assign, "assignments_matrix")
                
                df_all_pinned = extract_all_pinned_from_excel(uploaded_file_m2)
                if not df_all_pinned.empty: save_table_to_db(df_all_pinned, "pinned_slots")
                st.success(f"✅ Đã đóng băng {len(df_all_pinned)} tiết cố định!")
                st.rerun()

        edited_pinned = st.data_editor(load_table("pinned_slots"), num_rows="dynamic", use_container_width=True, height=400)
        if st.button("💾 LƯU KHUNG GIỜ", type="primary"): save_table_to_db(edited_pinned, "pinned_slots")

elif menu == "Module 3: Kiểm tra Lịch":
    st.markdown('<div class="main-title">🔍 MODULE 3: KIỂM TRA LỊCH CỐ ĐỊNH & NGUYÊN TẮC</div>', unsafe_allow_html=True)
    
    with st.expander("⚙️ BẢNG CẤU HÌNH NGUYÊN TẮC RÀNG BUỘC (RULES ENGINE)", expanded=True):
        st.info("💡 Hệ thống AI và Module kiểm tra sẽ quét tuân thủ theo các cài đặt trong bảng này. Bấm LƯU để áp dụng thiết lập mới.")
        if "rules_df" not in st.session_state:
            st.session_state.rules_df = load_rules_from_db()
            
        edited_rules = st.data_editor(st.session_state.rules_df, num_rows="dynamic", use_container_width=True, hide_index=True)
        st.session_state.rules_df = edited_rules
        
        if st.button("💾 LƯU BẢNG NGUYÊN TẮC", type="secondary"):
            save_rules_to_db_persist(edited_rules)
            st.success("Đã lưu bảng nguyên tắc vào hệ thống Database!")
            st.rerun()
            
        active_rules = {row['Mã Nguyên Tắc']: row['Bật/Tắt'] for _, row in edited_rules.iterrows()}

    with st.spinner("Đang tổng hợp ma trận và kiểm tra nguyên tắc..."):
        res1 = get_fixed_schedule(active_rules)
        if res1["status"] == "SUCCESS":
            if res1.get("warnings"):
                st.markdown("#### ⚠️ PHÁT HIỆN LỖI XUNG ĐỘT TRONG LỊCH CỐ ĐỊNH:")
                for w in res1["warnings"]:
                    st.warning(w)
            
            if "message" in res1 and res1["message"]:
                st.info(res1["message"])
                
            df_tkb = res1["data"]
            df_tkb['Ngày'] = pd.Categorical(df_tkb['Ngày'], categories=["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"], ordered=True)
            
            df_sang = df_tkb[df_tkb['Tiết'] <= 4]
            df_chieu = df_tkb[df_tkb['Tiết'] > 4]
            
            tab_s, tab_c = st.tabs(["🌞 Lịch Cố Định Sáng", "🌙 Lịch Cố Định Chiều"])
            with tab_s: st.dataframe(styler_func(pivot_main(df_sang)), use_container_width=True, height=600)
            with tab_c: st.dataframe(styler_func(pivot_main(df_chieu)), use_container_width=True, height=600)
            
            st.markdown("---")
            excel_bytes_m3 = generate_styled_excel_both_shifts(df_sang, df_chieu, df_tkb)
            st.download_button(
                label="📥 TẢI FILE EXCEL LỊCH CỐ ĐỊNH (Chuẩn Mẫu & Màu UI)",
                data=excel_bytes_m3,
                file_name="TKB_CoDinh_MauChuan.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True
            )
        else:
            st.error(res1["message"])

elif menu == "Module 4: AI Lấp Đầy":
    st.markdown('<div class="main-title">🧠 MODULE 4: AI LẤP ĐẦY THỜI KHÓA BIỂU</div>', unsafe_allow_html=True)
    
    active_rules = {row['Mã Nguyên Tắc']: row['Bật/Tắt'] for _, row in load_rules_from_db().iterrows()}

    st.markdown("### 📌 VÒNG 1: KHỞI TẠO KHUNG CỐ ĐỊNH")
    if st.button("🚀 CHẠY VÒNG 1: KHỞI TẠO KHUNG CỐ ĐỊNH", type="primary", use_container_width=True):
        with st.spinner("Đang lên khung TKB cốt lõi..."):
            res1 = run_round_1(active_rules)
            if res1["status"] == "SUCCESS":
                st.session_state.tkb_round1 = res1["data"]
                st.session_state.r1_vars = res1["r1_vars"]
                st.success("✅ Đã xử lý xong Vòng 1!")
            else:
                st.error(res1.get("message", "Lỗi không xác định ở Vòng 1"))

    if "tkb_round1" in st.session_state:
        df_tkb1 = st.session_state.tkb_round1.copy()
        df_tkb1['Ngày'] = pd.Categorical(df_tkb1['Ngày'], categories=["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"], ordered=True)
        tab_s1, tab_c1 = st.tabs(["🌞 Khung TKB Sáng (Vòng 1)", "🌙 Khung TKB Chiều (Vòng 1)"])
        with tab_s1: st.dataframe(styler_func(pivot_main(df_tkb1[df_tkb1['Tiết'] <= 4])), use_container_width=True, height=400)
        with tab_c1: st.dataframe(styler_func(pivot_main(df_tkb1[df_tkb1['Tiết'] > 4])), use_container_width=True, height=400)
        
        st.markdown("---")
        st.markdown("### 🌟 VÒNG 2: BỔ SUNG MÔN & BỊT KÍN SÁNG CƠ BẢN")
        
        if st.button("🚀 CHẠY VÒNG 2: ĐẮP MÔN BỔ SUNG", type="primary", use_container_width=True):
            with st.spinner("Đang tổng lực rải tiết... (Quá trình mất khoảng 15-25 giây)..."):
                res2 = run_round_2(st.session_state.r1_vars, active_rules)
                if res2["status"] == "SUCCESS":
                    st.session_state.tkb_round2 = res2["data"]
                    st.session_state.r2_vars = res2["r2_vars"]
                    st.success("✅ Đã hoàn thành Vòng 2!")
                else:
                    st.error(res2.get("message", "Lỗi không xác định ở Vòng 2"))

    if "tkb_round2" in st.session_state:
        st.markdown("#### 📝 CHỈNH SỬA MA TRẬN TKB SÁNG & CHIỀU (VÒNG 2):")
        st.info("💡 CÁCH 1: Bạn có thể trực tiếp chỉnh sửa các ô trống trong lưới TKB dưới đây, sau đó bấm nút LƯU MA TRẬN để cập nhật lại dữ liệu cho Vòng 3.")
        
        df_tkb2 = st.session_state.tkb_round2.copy()
        df_tkb2['Ngày'] = pd.Categorical(df_tkb2['Ngày'], categories=["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"], ordered=True)
        
        df_sang2 = df_tkb2[df_tkb2['Tiết'] <= 4]
        df_chieu2 = df_tkb2[df_tkb2['Tiết'] > 4]
        
        pt_sang = pivot_main(df_sang2)
        pt_chieu = pivot_main(df_chieu2)
        
        pt_sang_editable = pt_sang.reset_index()
        pt_chieu_editable = pt_chieu.reset_index()
        
        tab_es2, tab_ec2 = st.tabs(["🌞 Sửa TKB Sáng (Vòng 2)", "🌙 Sửa TKB Chiều (Vòng 2)"])
        
        with tab_es2:
            edited_sang_flat = st.data_editor(pt_sang_editable, use_container_width=True, height=400, key="edit_sang_flat_v2")
        with tab_ec2:
            edited_chieu_flat = st.data_editor(pt_chieu_editable, use_container_width=True, height=400, key="edit_chieu_flat_v2")
            
        if st.button("💾 LƯU MA TRẬN CHỈNH SỬA (SỬA TRỰC TIẾP TRÊN WEB)", type="primary"):
            new_rows = []
            def unpivot_flat(flat_df):
                for _, row in flat_df.iterrows():
                    d = row['Ngày']
                    p = row['Tiết']
                    for col in flat_df.columns:
                        if col not in ['Ngày', 'Tiết']:
                            val = str(row[col]).strip()
                            if val and val != 'nan':
                                new_rows.append({"Lớp": col, "Giáo viên": val, "Ngày": d, "Tiết": p})
                                
            unpivot_flat(edited_sang_flat)
            unpivot_flat(edited_chieu_flat)
            
            new_df_r2 = pd.DataFrame(new_rows)
            st.session_state.tkb_round2 = new_df_r2
            st.session_state.r2_vars = parse_r2_vars_from_df(new_df_r2)
            st.success("✅ Đã lưu ma trận chỉnh sửa thành công! Bạn có thể kéo xuống bấm CHẠY VÒNG 3 để kiểm tra lỗi.")

        st.markdown("---")
        st.info("💡 CÁCH 2: Tải file Excel ở nút bên dưới về. Chỉnh sửa tùy ý trên máy tính (bằng phần mềm Excel), sau đó upload ngược file vừa sửa vào ô bên cạnh để hệ thống đồng bộ!")
        
        col_down, col_up = st.columns([1, 1])
        with col_down:
            excel_bytes_m2 = generate_styled_excel_both_shifts(df_sang2, df_chieu2, st.session_state.tkb_round2)
            st.download_button(
                label="📥 BƯỚC 1: TẢI BẢN SAO LƯU EXCEL TKB (Để sửa tay)",
                data=excel_bytes_m2,
                file_name="TKB_KetQua_Vong2.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="secondary",
                use_container_width=True
            )
            
        with col_up:
            uploaded_tkb = st.file_uploader("📤 BƯỚC 2: NẠP LẠI FILE EXCEL ĐÃ SỬA TAY", type=["xlsx"], label_visibility="collapsed")
            if uploaded_tkb is not None:
                if st.button("🔄 Cập nhật dữ liệu từ file Excel (Đã sửa)", type="primary", use_container_width=True):
                    try:
                        xls = pd.ExcelFile(uploaded_tkb)
                        all_rows = []
                        for sheet in ["TKB Sáng", "TKB Chiều"]:
                            if sheet in xls.sheet_names:
                                df_sheet = pd.read_excel(xls, sheet_name=sheet)
                                if 'Ngày' in df_sheet.columns:
                                    df_sheet['Ngày'] = df_sheet['Ngày'].ffill() 
                                    for _, row in df_sheet.iterrows():
                                        d = row['Ngày']
                                        p = row['Tiết']
                                        for col in df_sheet.columns:
                                            if col not in ['Ngày', 'Tiết'] and not str(col).startswith('Unnamed'):
                                                val = str(row[col]).strip()
                                                if val and val != 'nan':
                                                    all_rows.append({"Lớp": col, "Giáo viên": val, "Ngày": d, "Tiết": p})
                        if all_rows:
                            new_df_r2 = pd.DataFrame(all_rows)
                            st.session_state.tkb_round2 = new_df_r2
                            st.session_state.r2_vars = parse_r2_vars_from_df(new_df_r2)
                            st.success("✅ Đã cập nhật ma trận từ file Excel thành công! Vui lòng bấm CHẠY VÒNG 3 ở bên dưới để kiểm tra lỗi và xuất bản.")
                            st.rerun()
                        else:
                            st.error("❌ Không tìm thấy dữ liệu hợp lệ. Đảm bảo file giữ nguyên sheet 'TKB Sáng' và 'TKB Chiều'.")
                    except Exception as e:
                        st.error(f"❌ Lỗi đọc file Excel: {e}")

        st.markdown("---")
        st.markdown("### 🏆 VÒNG 3: KIỂM TRA ĐỦ TIẾT & CHỐNG ĐÂM ĐỤNG")

        if st.button("🚀 CHẠY VÒNG 3: HOÀN THIỆN TOÀN DIỆN", type="primary", use_container_width=True):
            with st.spinner("Đang chạy kiểm toán chéo Lớp học và Giáo viên... (Có thể mất đến 30 giây)..."):
                res3 = run_round_3(st.session_state.r2_vars, active_rules)
                if res3["status"] == "SUCCESS":
                    st.session_state.final_tkb = res3["data"]
                    st.session_state.diagnosis = res3["diagnosis"]
                    st.session_state.audit_class = res3.get("audit_class", None)
                    st.session_state.audit_gv = res3.get("audit_gv", None)
                else:
                    st.error(res3.get("message", "Lỗi đâm đụng ở Vòng 3"))

    if "final_tkb" in st.session_state and st.session_state.final_tkb is not None:
        if "diagnosis" in st.session_state and len(st.session_state.diagnosis) > 0:
            st.markdown("#### ⚠️ BÁO CÁO KIỂM TOÁN TÌNH TRẠNG LỊCH:")
            def extract_class_for_sort(msg):
                m = re.search(r'(10|11|12)A\d+', msg)
                if m:
                    parts = re.match(r'^(\d+)A(\d+)$', m.group(0))
                    if parts: return (int(parts.group(1)), int(parts.group(2)), msg)
                return (99, 99, msg)
                
            sorted_diagnosis = sorted(st.session_state.diagnosis, key=extract_class_for_sort)
            for diag in sorted_diagnosis:
                if '🔴' in diag or '❌' in diag: st.error(diag) 
                elif '⚠️' in diag: st.warning(diag) 
                else: st.success(diag)
        else:
            st.success("✅ ĐÃ TẠO TKB THÀNH CÔNG 100% HOÀN MỸ, KHÔNG CÓ LỖI!")
            
        if "audit_class" in st.session_state and st.session_state.audit_class is not None:
            with st.expander("📊 BẢNG ĐỐI CHIẾU TIẾT (Module 1 vs Module 4)", expanded=False):
                tab_ac, tab_ag = st.tabs(["🏫 Đối chiếu theo Lớp", "👨‍🏫 Đối chiếu theo Giáo viên"])
                with tab_ac: st.dataframe(st.session_state.audit_class, use_container_width=True, height=400)
                with tab_ag: st.dataframe(st.session_state.audit_gv, use_container_width=True, height=400)
            
        df_tkb = st.session_state.final_tkb.copy()
        df_tkb['Ngày'] = pd.Categorical(df_tkb['Ngày'], categories=["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"], ordered=True)

        df_sang = df_tkb[df_tkb['Tiết'] <= 4]
        df_chieu = df_tkb[df_tkb['Tiết'] > 4]

        tab1, tab2, tab3, tab4 = st.tabs(["🌞 TKB Sáng", "🌙 TKB Chiều", "🏫 TKB Từng Lớp", "👨‍🏫 TKB Từng Giáo Viên"])
        
        with tab1: st.dataframe(styler_func(pivot_main(df_sang)), use_container_width=True, height=600)
        with tab2: st.dataframe(styler_func(pivot_main(df_chieu)), use_container_width=True, height=600)
            
        with tab3:
            col_lop = st.selectbox("📌 Chọn Lớp để xem chi tiết:", sorted(df_tkb['Lớp'].unique(), key=lambda x: (int(re.search(r'\d+', x).group()) if re.search(r'\d+', x) else 99, x)))
            if col_lop: 
                pt_lop = pivot_single(df_tkb[df_tkb['Lớp'] == col_lop], 'Giáo viên')
                st.table(pt_lop)
                
                excel_lop = generate_single_styled_excel(pt_lop, col_lop, "Lớp")
                st.download_button(
                    label=f"📥 TẢI EXCEL LỊCH HỌC LỚP {col_lop}",
                    data=excel_lop,
                    file_name=f"TKB_Lop_{col_lop}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )
                
        with tab4:
            all_gvs = set()
            for gv_str in df_tkb['Giáo viên'].unique():
                if '⭐' not in gv_str and '🟢' not in gv_str and '📌' not in gv_str and '⛔' not in gv_str:
                    clean_name = gv_str.replace('🔴', '').split('-')[-1].strip() if '-' in gv_str else gv_str.replace('🔴', '').strip()
                    all_gvs.add(clean_name)
                    
            selected_gv = st.selectbox("📌 Chọn Giáo viên để xem chi tiết:", sorted(list(all_gvs)))
            if selected_gv:
                df_gv = df_tkb[df_tkb['Giáo viên'].str.contains(selected_gv, na=False, regex=False)]
                pt_gv = pivot_single(df_gv, 'Lớp')
                st.table(pt_gv)
                
                excel_gv = generate_single_styled_excel(pt_gv, selected_gv, "Giáo viên")
                st.download_button(
                    label=f"📥 TẢI EXCEL LỊCH DẠY GIÁO VIÊN {selected_gv}",
                    data=excel_gv,
                    file_name=f"TKB_GV_{selected_gv}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )

        st.markdown("---")
        excel_bytes_final = generate_styled_excel_both_shifts(df_sang, df_chieu, df_tkb)
        st.download_button(
            label="📥 TẢI FILE EXCEL TKB HOÀN CHỈNH TỐI ƯU NHẤT (Màu Chuẩn UI)",
            data=excel_bytes_final,
            file_name="TKB_HoanChinh_Final.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True
        )